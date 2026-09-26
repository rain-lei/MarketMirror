from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .common import (
    canonical_headers,
    normalize_reply_status,
    normalize_stock_code,
    parse_datetime,
    row_looks_like_header,
)


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return value
    return str(value)


def _update_stats(stats: dict[str, Any], row: list[Any], headers: list[str]) -> None:
    stats["records"] += 1
    values = dict(zip(headers, row))
    for name, value in values.items():
        if value in (None, ""):
            stats["missing_by_field"][name] += 1

    code = normalize_stock_code(values.get("stock_code"))
    company = values.get("company_name")
    if code:
        stats["stock_codes"].add(code)
    if company not in (None, ""):
        stats["company_names"].add(str(company).strip())

    question_time = parse_datetime(values.get("question_time"))
    if question_time:
        current_min = stats["question_time_min"]
        current_max = stats["question_time_max"]
        if current_min is None or question_time < current_min:
            stats["question_time_min"] = question_time
        if current_max is None or question_time > current_max:
            stats["question_time_max"] = question_time

    status = normalize_reply_status(values.get("reply_status"))
    stats["reply_status"][status or "missing"] += 1
    user_type = values.get("user_type")
    stats["user_type"][str(user_type).strip() if user_type not in (None, "") else "missing"] += 1
    quarter = values.get("fiscal_quarter")
    if quarter not in (None, ""):
        stats["fiscal_quarter"][str(quarter)] += 1
    label = values.get("label_company_violation")
    if label not in (None, ""):
        stats["company_violation_label"][str(label)] += 1

    question = values.get("question_text")
    reply = values.get("reply_text")
    if question not in (None, ""):
        stats["question_length"]["n"] += 1
        stats["question_length"]["min"] = min(stats["question_length"]["min"], len(str(question)))
        stats["question_length"]["max"] = max(stats["question_length"]["max"], len(str(question)))
        stats["question_length"]["sum"] += len(str(question))
    if reply not in (None, ""):
        stats["reply_length"]["n"] += 1
        stats["reply_length"]["min"] = min(stats["reply_length"]["min"], len(str(reply)))
        stats["reply_length"]["max"] = max(stats["reply_length"]["max"], len(str(reply)))
        stats["reply_length"]["sum"] += len(str(reply))


def _new_stats() -> dict[str, Any]:
    return {
        "records": 0,
        "skipped_header_like_rows": 0,
        "blank_rows": 0,
        "missing_by_field": Counter(),
        "stock_codes": set(),
        "company_names": set(),
        "question_time_min": None,
        "question_time_max": None,
        "reply_status": Counter(),
        "user_type": Counter(),
        "fiscal_quarter": Counter(),
        "company_violation_label": Counter(),
        "question_length": {"n": 0, "min": 10**9, "max": 0, "sum": 0},
        "reply_length": {"n": 0, "min": 10**9, "max": 0, "sum": 0},
        "sample_rows": [],
    }


def _finalize_stats(stats: dict[str, Any]) -> dict[str, Any]:
    for field in ("question_length", "reply_length"):
        item = stats[field]
        if item["n"] == 0:
            stats[field] = None
        else:
            stats[field] = {
                "n": item["n"],
                "min": item["min"],
                "max": item["max"],
                "mean": round(item["sum"] / item["n"], 3),
            }
    for key in ("missing_by_field", "reply_status", "user_type", "fiscal_quarter", "company_violation_label"):
        stats[key] = dict(stats[key].most_common())
    stats["stock_code_count"] = len(stats.pop("stock_codes"))
    stats["company_name_count"] = len(stats.pop("company_names"))
    for key in ("question_time_min", "question_time_max"):
        if stats[key] is not None:
            stats[key] = stats[key].isoformat(sep=" ")
    return stats


def profile_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise ValueError("workbook has no worksheets")
        sheet = workbook.worksheets[0]
        rows: Iterable[tuple[Any, ...]] = sheet.iter_rows(values_only=True)
        first_rows: list[list[Any]] = []
        for _ in range(5):
            try:
                first_rows.append(list(next(rows)))
            except StopIteration:
                break
        if not first_rows:
            raise ValueError("worksheet is empty")

        header_index = 0
        header = first_rows[0]
        canonical = canonical_headers(header)
        stats = _new_stats()

        for index, candidate in enumerate(first_rows[1:], start=1):
            if row_looks_like_header(candidate, header):
                stats["skipped_header_like_rows"] += 1
                continue
            header_index = 0
            _update_stats(stats, candidate, canonical)
            if len(stats["sample_rows"]) < 5:
                stats["sample_rows"].append({name: _json_value(value) for name, value in zip(canonical, candidate)})

        for candidate in rows:
            if not any(value not in (None, "") for value in candidate):
                stats["blank_rows"] += 1
                continue
            if row_looks_like_header(list(candidate), header):
                stats["skipped_header_like_rows"] += 1
                continue
            _update_stats(stats, list(candidate), canonical)
            if len(stats["sample_rows"]) < 5:
                stats["sample_rows"].append({name: _json_value(value) for name, value in zip(canonical, candidate)})

        return {
            "file": path.name,
            "path": str(path),
            "size_bytes": path.stat().st_size,
            "sheet": sheet.title,
            "header_row": header_index + 1,
            "raw_headers": [_json_value(value) for value in header],
            "canonical_headers": canonical,
            "profile": _finalize_stats(stats),
        }
    finally:
        workbook.close()


def render_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# MarketMirror 数据质量报告",
        "",
        f"生成时间：{report['generated_at']}",
        "",
        "本报告由研究数据管道生成。原始文件不会被修改。",
        "",
    ]
    for item in report["sources"]:
        profile = item["profile"]
        lines.extend(
            [
                f"## {item['file']}",
                "",
                f"- 工作表：`{item['sheet']}`",
                f"- 字段数：{len(item['raw_headers'])}",
                f"- 数据行数：{profile['records']}",
                f"- 跳过的表头或释义行：{profile['skipped_header_like_rows']}",
                f"- 股票代码数：{profile['stock_code_count']}",
                f"- 公司名称数：{profile['company_name_count']}",
                f"- 提问时间：{profile['question_time_min']} 至 {profile['question_time_max']}",
                f"- 回复状态：`{json.dumps(profile['reply_status'], ensure_ascii=False)}`",
                f"- 用户类型：`{json.dumps(profile['user_type'], ensure_ascii=False)}`",
                f"- 缺失最多的字段：`{json.dumps(dict(list(profile['missing_by_field'].items())[:12]), ensure_ascii=False)}`",
                "",
                "### 规范化字段",
                "",
                "```text",
                ", ".join(item["canonical_headers"]),
                "```",
                "",
            ]
        )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Profile MarketMirror Excel sources")
    parser.add_argument("--output-dir", type=Path, default=Path("research_outputs"))
    parser.add_argument("sources", nargs="+", type=Path)
    args = parser.parse_args()

    from datetime import datetime, timezone

    args.output_dir.mkdir(parents=True, exist_ok=True)
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "sources": [profile_workbook(source) for source in args.sources],
    }
    (args.output_dir / "data_quality_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.output_dir / "data_quality_report.md").write_text(render_markdown(report), encoding="utf-8")
    print(json.dumps({"output_dir": str(args.output_dir), "sources": len(report["sources"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
