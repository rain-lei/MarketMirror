from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from .common import (
    canonical_headers,
    make_qa_id,
    normalize_reply_status,
    normalize_stock_code,
    iso_or_none,
    row_looks_like_header,
)


def normalize_row(row: list[object], headers: list[str], source_file: str) -> dict[str, object]:
    values = dict(zip(headers, row))
    code = normalize_stock_code(values.get("stock_code"))
    question_time = iso_or_none(values.get("question_time"))
    result: dict[str, object] = {
        "qa_id": make_qa_id(source_file, code, question_time, values.get("question_text")),
        "source": values.get("source"),
        "stock_code": code,
        "company_name": values.get("company_name"),
        "user_type": values.get("user_type"),
        "user_name": values.get("user_name"),
        "question_time": question_time,
        "question_text": values.get("question_text"),
        "reply_status": normalize_reply_status(values.get("reply_status")),
        "reply_time": iso_or_none(values.get("reply_time")),
        "reply_text": values.get("reply_text"),
        "question_length_source": values.get("question_length_source"),
        "reply_length_source": values.get("reply_length_source"),
        "label_company_violation": values.get("label_company_violation"),
        "fiscal_quarter": values.get("fiscal_quarter"),
    }
    for key, value in values.items():
        if key.startswith("financial__"):
            result[key] = value
    return result


def normalize_workbook(input_path: Path, output_path: Path, limit: int | None = None) -> dict[str, int]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_header = list(next(rows))
        except StopIteration as exc:
            raise ValueError("worksheet is empty") from exc
        headers = canonical_headers(raw_header)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        fieldnames = [
            "qa_id",
            "source",
            "stock_code",
            "company_name",
            "user_type",
            "user_name",
            "question_time",
            "question_text",
            "reply_status",
            "reply_time",
            "reply_text",
            "question_length_source",
            "reply_length_source",
            "label_company_violation",
            "fiscal_quarter",
        ] + [name for name in headers if name.startswith("financial__")]
        written = 0
        skipped = 0
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                values = list(row)
                if not any(value not in (None, "") for value in values):
                    skipped += 1
                    continue
                if row_looks_like_header(values, raw_header):
                    skipped += 1
                    continue
                writer.writerow(normalize_row(values, headers, input_path.name))
                written += 1
                if limit is not None and written >= limit:
                    break
        return {"written": written, "skipped": skipped}
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize a MarketMirror Q&A workbook to CSV")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args()
    print(normalize_workbook(args.input, args.output, args.limit))


if __name__ == "__main__":
    main()
