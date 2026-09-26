"""Build a traced multi-source Q&A dataset, keeping financial inputs quarantined."""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import tempfile
from collections import Counter
from contextlib import closing
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import canonical_headers, normalize_reply_status, row_looks_like_header
from .extract_financial import CATALOG_PATH, _as_text, _json, _raw_value
from .provenance import file_sha256

VERSION = "qa-dataset-v1"
SOURCE_TIMEZONE = timezone(timedelta(hours=8))
POLICY_PATH = Path(__file__).resolve().parents[1] / "data_contracts" / "forecast_feature_policy.json"

SCHEMA = """
CREATE TABLE source_file (
    source_file_hash TEXT PRIMARY KEY, source_path TEXT NOT NULL,
    source_sheet TEXT NOT NULL, header_json TEXT NOT NULL, row_count INTEGER NOT NULL
);
CREATE TABLE company (
    company_id TEXT PRIMARY KEY, stock_code TEXT NOT NULL UNIQUE,
    identity_basis TEXT NOT NULL CHECK(identity_basis = 'source_stock_code')
);
CREATE TABLE company_alias (
    company_id TEXT NOT NULL REFERENCES company(company_id), company_name TEXT NOT NULL,
    source_file_hash TEXT NOT NULL REFERENCES source_file(source_file_hash),
    reference_count INTEGER NOT NULL,
    PRIMARY KEY(company_id, company_name, source_file_hash)
);
CREATE TABLE financial_snapshot (
    snapshot_id TEXT PRIMARY KEY, source_file_hash TEXT NOT NULL REFERENCES source_file(source_file_hash),
    stock_code TEXT, source_quarter_hint TEXT, metric_values_json TEXT NOT NULL,
    report_period_end TEXT, published_at TEXT,
    verification_status TEXT NOT NULL CHECK(verification_status = 'unverified'), reference_count INTEGER NOT NULL
);
CREATE TABLE source_row_reference (
    source_file_hash TEXT NOT NULL REFERENCES source_file(source_file_hash), source_sheet TEXT NOT NULL,
    source_row INTEGER NOT NULL, snapshot_id TEXT NOT NULL REFERENCES financial_snapshot(snapshot_id),
    company_name TEXT, question_time_raw TEXT, reply_time_raw TEXT, violation_label_raw TEXT,
    context_values_json TEXT NOT NULL,
    PRIMARY KEY(source_file_hash, source_sheet, source_row)
);
CREATE TABLE financial_field_dictionary (
    source_file_hash TEXT NOT NULL REFERENCES source_file(source_file_hash), source_column INTEGER NOT NULL,
    field_name TEXT NOT NULL, role TEXT NOT NULL, details_json TEXT NOT NULL,
    PRIMARY KEY(source_file_hash, source_column)
);
CREATE TABLE qa_record (
    qa_id TEXT PRIMARY KEY, content_id TEXT NOT NULL,
    source_file_hash TEXT NOT NULL REFERENCES source_file(source_file_hash),
    source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
    company_id TEXT REFERENCES company(company_id), stock_code TEXT, company_name TEXT,
    source TEXT, user_type TEXT, user_name TEXT,
    question_available_at TEXT, question_text TEXT, question_eligible INTEGER NOT NULL CHECK(question_eligible IN (0,1)),
    reply_available_at TEXT, reply_text TEXT,
    reply_eligible INTEGER NOT NULL CHECK(reply_eligible IN (0,1)),
    snapshot_id TEXT REFERENCES financial_snapshot(snapshot_id), quality_flags_json TEXT NOT NULL,
    UNIQUE(source_file_hash, source_sheet, source_row)
);
CREATE TABLE qa_restricted_context (
    qa_id TEXT PRIMARY KEY REFERENCES qa_record(qa_id), raw_metadata_json TEXT NOT NULL,
    extra_fields_json TEXT NOT NULL
);
CREATE TABLE dataset_metadata (key TEXT PRIMARY KEY, value_json TEXT NOT NULL);
CREATE INDEX qa_question_time ON qa_record(question_available_at, stock_code) WHERE question_eligible = 1;
CREATE INDEX qa_content ON qa_record(content_id);
CREATE INDEX row_snapshot ON source_row_reference(snapshot_id);
"""


def normalize_timestamp(value: Any) -> str | None:
    """Convert explicit timestamps to fixed-width +08:00, rejecting date-only values."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date) or value in (None, ""):
        return None
    else:
        text = str(value).strip().replace("/", "-")
        if "T" not in text and " " not in text:
            return None
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError:
            return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=SOURCE_TIMEZONE)
    return parsed.astimezone(SOURCE_TIMEZONE).isoformat(timespec="microseconds")


def valid_stock_code(value: Any) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    match = re.fullmatch(r"([0-9]{1,6})(?:\.0+)?", str(value).strip())
    return match.group(1).zfill(6) if match else None


def normalize_reply_timestamp(value: Any) -> tuple[str | None, bool]:
    """For a date-only reply, use the following midnight as a conservative bound."""
    timestamp = normalize_timestamp(value)
    if timestamp is not None:
        return timestamp, False
    if isinstance(value, date) and not isinstance(value, datetime):
        day = value
    elif isinstance(value, str) and re.fullmatch(r"[0-9]{4}[-/][0-9]{2}[-/][0-9]{2}", value.strip()):
        try:
            day = date.fromisoformat(value.strip().replace("/", "-"))
        except ValueError:
            return None, False
    else:
        return None, False
    try:
        bound = datetime.combine(day + timedelta(days=1), datetime.min.time(), tzinfo=SOURCE_TIMEZONE)
    except OverflowError:
        return None, False
    return bound.isoformat(timespec="microseconds"), True


def _digest(value: Any) -> str:
    return hashlib.sha256(_json(value).encode("utf-8")).hexdigest()


def _usable_text(value: Any) -> bool:
    return (isinstance(value, str) and bool(value.strip()) and not value.startswith("=")
            and value not in {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"})


def _financial_sources(paths: list[Path]) -> dict[str, dict[str, Any]]:
    result = {}
    for path in paths:
        path = path.resolve()
        digest = file_sha256(path)
        with closing(sqlite3.connect(path.as_uri() + "?mode=ro", uri=True)) as db:
            metadata = {key: json.loads(value) for key, value in db.execute("SELECT key, value_json FROM run_metadata")}
            if metadata.get("pipeline_version") != "financial-extract-v1":
                raise ValueError("unsupported financial extraction version")
            source_hash = metadata["source_file_hash"]
            if source_hash in result:
                raise ValueError("multiple financial databases for one source file")
            if db.execute("PRAGMA integrity_check").fetchone()[0] != "ok" or db.execute("PRAGMA foreign_key_check").fetchall():
                raise ValueError("financial database integrity check failed")
            actual = db.execute("SELECT COUNT(*) FROM source_row_reference").fetchone()[0]
            if actual != metadata["counts"]["source_rows"]:
                raise ValueError("financial reference count disagrees with metadata")
        result[source_hash] = {"path": path, "sha256": digest, "metadata": metadata}
    return result


def _import_financial(conn: sqlite3.Connection, source_hash: str, info: dict[str, Any]) -> None:
    conn.execute("ATTACH DATABASE ? AS fin", (info["path"].as_uri() + "?mode=ro",))
    try:
        wrong = conn.execute(
            "SELECT COUNT(*) FROM fin.financial_snapshot WHERE source_file_hash != ? OR verification_status != 'unverified' "
            "OR published_at IS NOT NULL OR report_period_end IS NOT NULL", (source_hash,)).fetchone()[0]
        if wrong:
            raise ValueError("financial snapshots must match the source and remain unverified")
        conn.execute("INSERT INTO financial_snapshot SELECT * FROM fin.financial_snapshot")
        conn.execute(
            "INSERT INTO source_row_reference SELECT ?, source_sheet, source_row, snapshot_id, company_name, "
            "question_time_raw, reply_time_raw, violation_label_raw, context_values_json FROM fin.source_row_reference",
            (source_hash,))
        cursor = conn.execute("SELECT * FROM fin.field_dictionary")
        names = [item[0] for item in cursor.description]
        for row in cursor.fetchall():
            field = dict(zip(names, row))
            conn.execute("INSERT INTO financial_field_dictionary VALUES (?, ?, ?, ?, ?)",
                         (source_hash, field["source_column"], field["field_name"], field["role"], _json(field)))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.execute("DETACH DATABASE fin")


def build_dataset(sources: list[Path], output_dir: Path, financial_databases: list[Path] | None = None) -> dict[str, Any]:
    """Stream source occurrences into a fresh database; atomically publish after validation."""
    if not sources:
        raise ValueError("at least one source workbook is required")
    inputs = [{"path": p.resolve(), "sha256": file_sha256(p)} for p in sources]
    source_hashes = [item["sha256"] for item in inputs]
    if len(set(source_hashes)) != len(inputs):
        raise ValueError("duplicate source file content; each workbook must be supplied once")
    finances = _financial_sources(financial_databases or [])
    if set(finances) - set(source_hashes):
        raise ValueError("financial database source hash does not match any supplied workbook")
    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    financial_names = set(catalog["snapshot_metrics"] + catalog["row_context_metrics"])
    policy_hash = file_sha256(POLICY_PATH)
    code_hashes = {name: file_sha256(Path(__file__).with_name(name)) for name in (
        "build_dataset.py", "common.py", "extract_financial.py", "provenance.py")}
    dataset_id = _digest([VERSION, sorted(source_hashes), sorted(v["sha256"] for v in finances.values()),
                          policy_hash, code_hashes, file_sha256(CATALOG_PATH)])
    output_dir = output_dir.resolve()
    protected = {item["path"] for item in inputs} | {v["path"] for v in finances.values()} | {CATALOG_PATH, POLICY_PATH}
    if output_dir / "dataset.sqlite" in protected or output_dir / "run_manifest.json" in protected:
        raise ValueError("output must not overwrite a source file")
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_dir, suffix=".sqlite.tmp", delete=False) as handle:
        temp_path = Path(handle.name)
    conn = sqlite3.connect(temp_path, uri=True)
    counts = Counter()
    source_reports = []
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.executescript(SCHEMA)
        for item in inputs:
            path, source_hash = item["path"], item["sha256"]
            book = load_workbook(path, read_only=True, data_only=False)
            local = Counter()
            try:
                sheet = book.worksheets[0]
                rows = sheet.iter_rows(values_only=True)
                raw_header = list(next(rows))
                headers = canonical_headers(raw_header)
                required = {"stock_code", "question_time", "question_text"}
                if not required <= set(headers):
                    raise ValueError(f"missing required QA headers in {path.name}")
                if len([h for h in headers if re.fullmatch(r"(?:stock_code|question_time|question_text)_\d+", h)]):
                    raise ValueError("duplicate required QA headers")
                has_financial = any(str(h).strip() in catalog["snapshot_metrics"] for h in raw_header)
                if has_financial and source_hash not in finances:
                    raise ValueError("financial workbook requires its matching extracted financial database")
                conn.execute("INSERT INTO source_file VALUES (?, ?, ?, ?, 0)",
                             (source_hash, str(path), sheet.title, _json([_raw_value(h) for h in raw_header])))
                if source_hash in finances:
                    _import_financial(conn, source_hash, finances[source_hash])
                    if finances[source_hash]["metadata"]["source_sheet"] != sheet.title:
                        raise ValueError("financial source worksheet differs from QA worksheet")
                for source_row, raw in enumerate(rows, start=2):
                    if not any(v not in (None, "") for v in raw):
                        local["blank_rows"] += 1
                        continue
                    if row_looks_like_header(list(raw), raw_header):
                        local["header_like_rows"] += 1
                        continue
                    values = dict(zip(headers, raw))
                    flags = []
                    code = valid_stock_code(values.get("stock_code"))
                    if code is None:
                        flags.append("invalid_stock_code")
                    company_id = f"cn-stock-code:{code}" if code else None
                    name = _as_text(values.get("company_name"))
                    if company_id:
                        conn.execute("INSERT OR IGNORE INTO company VALUES (?, ?, 'source_stock_code')", (company_id, code))
                        if name and name.strip():
                            conn.execute("INSERT INTO company_alias VALUES (?, ?, ?, 1) ON CONFLICT "
                                         "DO UPDATE SET reference_count = reference_count + 1", (company_id, name, source_hash))
                    question_at = normalize_timestamp(values.get("question_time"))
                    if question_at is None:
                        flags.append("invalid_or_date_only_question_time")
                    question = _as_text(values.get("question_text"))
                    if not _usable_text(values.get("question_text")):
                        flags.append("unusable_question_text")
                    eligible = not flags
                    reply = _as_text(values.get("reply_text"))
                    reply_at, date_only_reply = normalize_reply_timestamp(values.get("reply_time"))
                    if date_only_reply:
                        flags.append("reply_date_only_next_day_bound")
                    status = normalize_reply_status(values.get("reply_status"))
                    reply_eligible = bool(eligible and status == "replied" and _usable_text(values.get("reply_text"))
                                          and reply_at and (reply_at > question_at if date_only_reply else reply_at >= question_at))
                    if status == "replied" and not reply_eligible:
                        flags.append("reply_unusable_or_invalid_time")
                    if status != "replied" and (reply_at or reply):
                        flags.append("reply_metadata_on_unconfirmed_reply")
                    snapshot_id = None
                    if source_hash in finances:
                        ref = conn.execute(
                            "SELECT r.snapshot_id, s.stock_code, r.company_name, r.question_time_raw, r.reply_time_raw, "
                            "r.violation_label_raw FROM source_row_reference r JOIN financial_snapshot s USING(snapshot_id) "
                            "WHERE r.source_file_hash = ? AND r.source_sheet = ? AND r.source_row = ?",
                            (source_hash, sheet.title, source_row)).fetchone()
                        expected = (code, name, _as_text(values.get("question_time")),
                                    _as_text(values.get("reply_time")), _as_text(values.get("label_company_violation")))
                        if ref is None or ref[1:] != expected:
                            raise ValueError(f"financial reference mismatch at source row {source_row}")
                        snapshot_id = ref[0]
                        local["linked_financial_rows"] += 1
                    qa_id = _digest([source_hash, sheet.title, source_row])
                    content_id = _digest([code, question_at, question, values.get("user_type"), values.get("user_name")])
                    conn.execute("INSERT INTO qa_record VALUES (" + ",".join(["?"] * 19) + ")", (
                        qa_id, content_id, source_hash, sheet.title, source_row, company_id, code, name,
                        _as_text(values.get("source")), _as_text(values.get("user_type")), _as_text(values.get("user_name")),
                        question_at, question, int(eligible), reply_at, reply, int(reply_eligible), snapshot_id, _json(flags)
                    ))
                    raw_metadata = {k: _raw_value(v) for k, v in values.items()
                                    if not k.startswith(("financial__", "unnamed_")) and k not in (
                                        "question_text", "reply_text", "company_name", "source", "user_name", "user_type")}
                    extra = {str(h): _raw_value(v) for h, v, k in zip(raw_header, raw, headers)
                             if k.startswith(("financial__", "unnamed_")) and not (
                                 source_hash in finances and str(h).strip() in financial_names)}
                    conn.execute("INSERT INTO qa_restricted_context VALUES (?, ?, ?)", (qa_id, _json(raw_metadata), _json(extra)))
                    local["qa_rows"] += 1
                    local["eligible_question_rows"] += int(eligible)
                    local["eligible_reply_rows"] += int(reply_eligible)
                    for flag in flags:
                        local[flag] += 1
                    if local["qa_rows"] % 100000 == 0:
                        conn.commit()
                        print(f"{path.name}: processed {local['qa_rows']:,} QA rows", flush=True)
                if source_hash in finances:
                    expected_rows = finances[source_hash]["metadata"]["counts"]["source_rows"]
                    if local["linked_financial_rows"] != expected_rows:
                        raise ValueError("financial references do not reconcile with QA rows")
                conn.execute("UPDATE source_file SET row_count = ? WHERE source_file_hash = ?", (local["qa_rows"], source_hash))
                counts.update(local)
                source_reports.append({"path": str(path), "sha256": source_hash, "sheet": sheet.title, "counts": dict(local)})
            finally:
                book.close()
        for item in inputs:
            if file_sha256(item["path"]) != item["sha256"]:
                raise RuntimeError("source workbook changed during dataset build")
        for info in finances.values():
            if file_sha256(info["path"]) != info["sha256"]:
                raise RuntimeError("financial input changed during dataset build")
        if counts["qa_rows"] != conn.execute("SELECT COUNT(*) FROM qa_record").fetchone()[0]:
            raise RuntimeError("QA row counts do not reconcile")
        counts["companies"] = conn.execute("SELECT COUNT(*) FROM company").fetchone()[0]
        counts["company_aliases"] = conn.execute("SELECT COUNT(*) FROM company_alias").fetchone()[0]
        counts["financial_snapshots"] = conn.execute("SELECT COUNT(*) FROM financial_snapshot").fetchone()[0]
        counts["duplicate_content_excess_rows"] = conn.execute(
            "SELECT COALESCE(SUM(n - 1), 0) FROM (SELECT COUNT(*) n FROM qa_record GROUP BY content_id HAVING COUNT(*) > 1)"
        ).fetchone()[0]
        manifest = {
            "pipeline_version": VERSION, "dataset_id": dataset_id,
            "generated_at": datetime.now(timezone.utc).isoformat(), "sources": source_reports,
            "financial_inputs": [{"path": str(v["path"]), "sha256": v["sha256"]} for v in finances.values()],
            "code_sha256": code_hashes, "financial_catalog_sha256": file_sha256(CATALOG_PATH),
            "feature_policy_sha256": policy_hash, "counts": dict(counts),
            "assumptions": [
                "Naive source timestamps are interpreted as Asia/Shanghai (+08:00). Date-only questions are ineligible.",
                "Date-only replies use the next day at 00:00 +08:00 as a conservative availability bound, not an observed intraday time.",
                "Question and reply timestamps proxy public availability; publication logs and revision history are absent.",
                "Company identity uses source stock code only; exchange and legal entity identity remain unverified.",
                "QA rows are source occurrences. Content duplicates are audited and retained, without automatic deduplication.",
                "Financial values, violation labels, source quarters and row context are excluded from forecast features.",
            ],
        }
        for key, value in manifest.items():
            conn.execute("INSERT INTO dataset_metadata VALUES (?, ?)", (key, _json(value)))
        conn.commit()
        if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok" or conn.execute("PRAGMA foreign_key_check").fetchall():
            raise RuntimeError("unified dataset integrity check failed")
        conn.close()
        conn = None
        target = output_dir / "dataset.sqlite"
        temp_path.replace(target)
        manifest["artifacts"] = {"dataset.sqlite": {"sha256": file_sha256(target), "bytes": target.stat().st_size}}
        (output_dir / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return manifest
    finally:
        if conn is not None:
            conn.close()
        temp_path.unlink(missing_ok=True)
