"""Stream repeated financial values into a traced, unverified SQLite dataset.

The source quarter is a hint. Report year, units and publication dates remain
unknown, so this extraction never promotes values into financial_quarter.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sqlite3
import tempfile
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import canonical_headers, normalize_stock_code, parse_datetime, row_looks_like_header
from .provenance import file_sha256

CATALOG_PATH = Path(__file__).resolve().parents[1] / "data_contracts" / "financial_fields.json"

SQL_SCHEMA = """
CREATE TABLE financial_snapshot (
    snapshot_id TEXT PRIMARY KEY,
    source_file_hash TEXT NOT NULL,
    stock_code TEXT,
    source_quarter_hint TEXT,
    metric_values_json TEXT NOT NULL,
    report_period_end TEXT,
    published_at TEXT,
    verification_status TEXT NOT NULL CHECK (verification_status = 'unverified'),
    reference_count INTEGER NOT NULL
);
CREATE TABLE source_row_reference (
    source_row INTEGER PRIMARY KEY,
    snapshot_id TEXT NOT NULL REFERENCES financial_snapshot(snapshot_id),
    source_sheet TEXT NOT NULL,
    company_name TEXT,
    question_time_raw TEXT,
    reply_time_raw TEXT,
    violation_label_raw TEXT,
    context_values_json TEXT NOT NULL
);
CREATE TABLE field_dictionary (
    source_column INTEGER PRIMARY KEY,
    field_name TEXT NOT NULL,
    role TEXT NOT NULL,
    missing_count INTEGER NOT NULL,
    numeric_count INTEGER NOT NULL,
    non_numeric_count INTEGER NOT NULL,
    negative_count INTEGER NOT NULL,
    zero_count INTEGER NOT NULL,
    numeric_min REAL,
    numeric_max REAL,
    unit TEXT,
    transform TEXT,
    verification_status TEXT NOT NULL
);
CREATE TABLE run_metadata (key TEXT PRIMARY KEY, value_json TEXT NOT NULL);
CREATE INDEX snapshot_company_quarter ON financial_snapshot(stock_code, source_quarter_hint);
CREATE INDEX row_snapshot ON source_row_reference(snapshot_id);
"""


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _raw_value(value: Any) -> Any:
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return int(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def _as_text(value: Any) -> str | None:
    return str(value) if value is not None else None


def _field_stats() -> dict[str, Any]:
    return {
        "missing_count": 0, "numeric_count": 0, "non_numeric_count": 0,
        "negative_count": 0, "zero_count": 0, "numeric_min": None, "numeric_max": None,
    }


def _add_stat(stats: dict[str, Any], value: Any) -> None:
    if value is None or value == "":
        stats["missing_count"] += 1
    elif isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        stats["numeric_count"] += 1
        stats["negative_count"] += int(value < 0)
        stats["zero_count"] += int(value == 0)
        stats["numeric_min"] = value if stats["numeric_min"] is None else min(value, stats["numeric_min"])
        stats["numeric_max"] = value if stats["numeric_max"] is None else max(value, stats["numeric_max"])
    else:
        stats["non_numeric_count"] += 1


def _markdown(report: dict[str, Any]) -> str:
    c = report["counts"]
    lines = [
        "# 财务提取与字段核对报告", "",
        f"来源：{report['source_file']}",
        f"来源 SHA-256：`{report['source_file_hash']}`", "",
        f"读取 {c['source_rows']:,} 条原始数据，提取 {c['snapshot_count']:,} 份不同的财务快照。",
        f"同一股票代码和来源季度出现多个财务版本的分组数：{c['conflicting_company_quarter_groups']:,}。",
        f"提问季度与来源季度不一致的记录数：{c['question_quarter_mismatches']:,}。", "",
        "所有快照均为 unverified。报告期、公告时间、单位和数值变换尚未确认。",
        "`p`、`p/e` 保存在来源行的 context_values_json 中，不参与季度财务快照去重。",
        "快照的 reference_count 表示原问答表重复引用次数，不代表独立财务观测数。", "",
        "## 字段统计", "",
        "| 字段 | 层级 | 缺失 | 数值 | 非数值 | 负值 | 零值 | 最小值 | 最大值 |",
        "|---|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for item in report["fields"]:
        lines.append(
            f"| {item['field_name'].replace('|', '/')} | {item['role']} | {item['missing_count']} | "
            f"{item['numeric_count']} | {item['non_numeric_count']} | {item['negative_count']} | "
            f"{item['zero_count']} | {item['numeric_min']} | {item['numeric_max']} |"
        )
    lines += ["", "## 待核对", ""] + [f"- {warning}" for warning in report["warnings"]]
    return "\n".join(lines) + "\n"


def extract_financial(input_path: Path, output_dir: Path) -> dict[str, Any]:
    input_path = input_path.resolve()
    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    roles = {name: "snapshot" for name in catalog["snapshot_metrics"]}
    roles.update({name: "row_context" for name in catalog["row_context_metrics"]})
    source_hash = file_sha256(input_path)
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    conn = None
    temp_path = None
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        raw_header = list(next(rows))
        header_names = [str(v).strip() if v is not None else "" for v in raw_header]
        financial_columns = [(i, name, roles[name]) for i, name in enumerate(header_names) if name in roles]
        if not any(role == "snapshot" for _, _, role in financial_columns):
            raise ValueError("source contains no known financial snapshot fields")
        financial_names = [name for _, name, _ in financial_columns]
        if len(financial_names) != len(set(financial_names)):
            raise ValueError("duplicate financial headers must be resolved before extraction")
        headers = canonical_headers(raw_header)
        if "stock_code" not in headers:
            raise ValueError("stock code header is missing")
        indexes = {name: i for i, name in enumerate(headers)}
        output_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=output_dir, suffix=".sqlite.tmp", delete=False) as handle:
            temp_path = Path(handle.name)
        conn = sqlite3.connect(temp_path)
        conn.execute("PRAGMA foreign_keys = ON")
        conn.executescript(SQL_SCHEMA)
        stats = {name: _field_stats() for _, name, _ in financial_columns}
        counts = Counter()
        for source_row, raw in enumerate(rows, start=2):
            if not any(v not in (None, "") for v in raw):
                counts["blank_rows"] += 1
                continue
            if row_looks_like_header(list(raw), raw_header):
                counts["skipped_header_like_rows"] += 1
                continue
            counts["source_rows"] += 1
            def val(key: str) -> Any:
                i = indexes.get(key)
                return raw[i] if i is not None and i < len(raw) else None
            code = normalize_stock_code(val("stock_code"))
            if code is None or re.fullmatch(r"\d{6}", code) is None:
                counts["invalid_stock_codes"] += 1
                code = None
            quarter = _as_text(_raw_value(val("fiscal_quarter")))
            if quarter not in ("1", "2", "3", "4"):
                counts["invalid_quarter_hints"] += 1
            question_time = parse_datetime(val("question_time"))
            if question_time and quarter in ("1", "2", "3", "4"):
                counts["question_quarter_mismatches"] += int(str((question_time.month - 1) // 3 + 1) != quarter)
            snapshot_values = {}
            context_values = {}
            for i, name, role in financial_columns:
                value = raw[i] if i < len(raw) else None
                _add_stat(stats[name], value)
                (snapshot_values if role == "snapshot" else context_values)[name] = _raw_value(value)
            # Missing identifiers never collapse unrelated records into one snapshot.
            row_salt = source_row if code is None or quarter not in ("1", "2", "3", "4") else None
            payload = _json([source_hash, sheet.title, code, quarter, snapshot_values, row_salt])
            snapshot_id = hashlib.sha256(payload.encode("utf-8")).hexdigest()
            conn.execute(
                "INSERT INTO financial_snapshot VALUES (?, ?, ?, ?, ?, NULL, NULL, 'unverified', 1) "
                "ON CONFLICT(snapshot_id) DO UPDATE SET reference_count = reference_count + 1",
                (snapshot_id, source_hash, code, quarter, _json(snapshot_values)),
            )
            conn.execute(
                "INSERT INTO source_row_reference VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (source_row, snapshot_id, sheet.title, _as_text(val("company_name")),
                 _as_text(val("question_time")), _as_text(val("reply_time")),
                 _as_text(val("label_company_violation")), _json(context_values)),
            )
            if counts["source_rows"] % 100000 == 0:
                conn.commit()
                print(f"Processed {counts['source_rows']:,} source rows", flush=True)
        fields = []
        for i, name, role in financial_columns:
            field_info = {"field_name": name, "role": role, **stats[name],
                          "unit": None, "transform": None, "verification_status": "unverified"}
            fields.append(field_info)
            conn.execute(
                "INSERT INTO field_dictionary VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, NULL, 'unverified')",
                (i + 1, name, role, stats[name]["missing_count"], stats[name]["numeric_count"],
                 stats[name]["non_numeric_count"], stats[name]["negative_count"], stats[name]["zero_count"],
                 stats[name]["numeric_min"], stats[name]["numeric_max"]),
            )
        counts["snapshot_count"] = conn.execute("SELECT COUNT(*) FROM financial_snapshot").fetchone()[0]
        counts["conflicting_company_quarter_groups"] = conn.execute(
            "SELECT COUNT(*) FROM (SELECT stock_code, source_quarter_hint FROM financial_snapshot "
            "WHERE stock_code IS NOT NULL AND source_quarter_hint IN ('1','2','3','4') "
            "GROUP BY stock_code, source_quarter_hint HAVING COUNT(*) > 1)"
        ).fetchone()[0]
        reference_sum = conn.execute("SELECT SUM(reference_count) FROM financial_snapshot").fetchone()[0] or 0
        if reference_sum != counts["source_rows"]:
            raise RuntimeError("snapshot reference counts do not reconcile with source rows")
        if counts["source_rows"] != conn.execute("SELECT COUNT(*) FROM source_row_reference").fetchone()[0]:
            raise RuntimeError("source row references do not reconcile")
        if file_sha256(input_path) != source_hash:
            raise RuntimeError("source file changed during extraction")
        counts = {key: counts[key] for key in (
            "source_rows", "snapshot_count", "conflicting_company_quarter_groups", "question_quarter_mismatches",
            "invalid_stock_codes", "invalid_quarter_hints", "blank_rows", "skipped_header_like_rows")}
        report = {
            "pipeline_version": "financial-extract-v1", "generated_at": datetime.now(timezone.utc).isoformat(),
            "source_file": input_path.name, "source_file_hash": source_hash, "source_sheet": sheet.title,
            "catalog_sha256": file_sha256(CATALOG_PATH), "counts": counts, "fields": fields,
            "warnings": [
                "财务年度、来源季度含义和公告时间未确认；不能按提问年度推断财报日期。",
                "单位和数值变换未确认；负资产等值需核对口径，当前原值不会用于计算财务比率。",
                "p 和 p/e 的观测时间与计算口径未确认，仅作为来源行上下文保留。",
                "同一公司和来源季度有多个快照时保留全部版本，不自动选择、求平均或覆盖。",
                "Isviolated 是来源行的公司标签，时间含义未确认，禁止直接作为预测输入。",
            ],
        }
        metadata = {key: value for key, value in report.items() if key != "fields"}
        metadata["source_path"] = str(input_path)
        for key, value in metadata.items():
            conn.execute("INSERT INTO run_metadata VALUES (?, ?)", (key, _json(value)))
        conn.commit()
        if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("SQLite integrity check failed")
        conn.close()
        conn = None
        target = output_dir / "financial_unverified.sqlite"
        temp_path.replace(target)
        temp_path = None
        report["database_sha256"] = file_sha256(target)
        (output_dir / "financial_quality_report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        (output_dir / "financial_quality_report.md").write_text(_markdown(report), encoding="utf-8")
        return report
    finally:
        workbook.close()
        if conn is not None:
            conn.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract unverified financial snapshots with row-level provenance")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    report = extract_financial(args.input, args.output_dir)
    print(json.dumps(report["counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
