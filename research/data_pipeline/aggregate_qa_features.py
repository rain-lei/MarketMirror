from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import canonical_headers, normalize_reply_status, normalize_stock_code, parse_datetime, row_looks_like_header


KEYWORDS: dict[str, tuple[str, ...]] = {
    "earnings": ("业绩", "利润", "营收", "收入", "盈利"),
    "policy": ("政策", "新规", "监管", "合规", "证监会", "问询"),
    "risk": ("风险", "不确定", "预警", "危机", "压力"),
    "liquidity": ("现金流", "流动性", "债务", "负债", "融资", "偿债"),
    "capital_return": ("分红", "回购", "减持", "股息", "股东回报"),
    "pandemic": ("疫情", "新冠", "病毒", "封锁"),
    "governance": ("违规", "诉讼", "内控", "治理", "审计", "披露"),
}


@dataclass
class NumericStats:
    count: int = 0
    total: int = 0
    minimum: int | None = None
    maximum: int | None = None

    def add(self, value: Any) -> None:
        if value in (None, ""):
            return
        length = len(str(value))
        self.count += 1
        self.total += length
        self.minimum = length if self.minimum is None else min(self.minimum, length)
        self.maximum = length if self.maximum is None else max(self.maximum, length)

    def mean(self) -> float | None:
        return round(self.total / self.count, 3) if self.count else None


@dataclass
class FeatureAccumulator:
    stock_code: str | None = None
    company_name: str | None = None
    question_count: int = 0
    user_count: set[str] = field(default_factory=set)
    user_type: Counter[str] = field(default_factory=Counter)
    question_length: NumericStats = field(default_factory=NumericStats)
    keyword_count: Counter[str] = field(default_factory=Counter)
    fiscal_quarter: Counter[str] = field(default_factory=Counter)
    reply_status_post_event: Counter[str] = field(default_factory=Counter)
    violation_label: set[str] = field(default_factory=set)

    def add(self, values: dict[str, Any]) -> None:
        self.question_count += 1
        self.stock_code = self.stock_code or normalize_stock_code(values.get("stock_code"))
        if values.get("company_name") not in (None, ""):
            self.company_name = str(values["company_name"]).strip()
        user_name = values.get("user_name")
        if user_name not in (None, ""):
            self.user_count.add(str(user_name).strip())
        user_type = values.get("user_type")
        self.user_type[str(user_type).strip() if user_type not in (None, "") else "missing"] += 1
        self.question_length.add(values.get("question_text"))
        text = str(values.get("question_text") or "")
        for category, words in KEYWORDS.items():
            if any(word in text for word in words):
                self.keyword_count[category] += 1
        if values.get("fiscal_quarter") not in (None, ""):
            self.fiscal_quarter[str(values["fiscal_quarter"])] += 1
        status = normalize_reply_status(values.get("reply_status"))
        self.reply_status_post_event[status or "missing"] += 1
        if values.get("label_company_violation") not in (None, ""):
            self.violation_label.add(str(values["label_company_violation"]))

    def as_row(self, period: str) -> dict[str, Any]:
        row: dict[str, Any] = {
            "period": period,
            "stock_code": self.stock_code,
            "company_name": self.company_name,
            "question_count": self.question_count,
            "unique_user_count": len(self.user_count),
            "registered_user_count": self.user_type.get("注册用户", 0),
            "browser_user_count": self.user_type.get("浏览用户", 0),
            "missing_user_type_count": self.user_type.get("missing", 0),
            "question_length_mean": self.question_length.mean(),
            "question_length_min": self.question_length.minimum,
            "question_length_max": self.question_length.maximum,
            "fiscal_quarter_mode": self.fiscal_quarter.most_common(1)[0][0] if self.fiscal_quarter else None,
            "label_company_violation": next(iter(self.violation_label)) if len(self.violation_label) == 1 else None,
            "label_company_violation_consistent": len(self.violation_label) <= 1,
            "post_event_replied_count": self.reply_status_post_event.get("replied", 0),
            "post_event_unreplied_count": self.reply_status_post_event.get("unreplied", 0),
        }
        for category in KEYWORDS:
            count = self.keyword_count.get(category, 0)
            row[f"{category}_count"] = count
            row[f"{category}_share"] = round(count / self.question_count, 6) if self.question_count else 0.0
        return row


def calendar_period(value: Any) -> str | None:
    parsed = parse_datetime(value)
    if parsed is None:
        return None
    return f"{parsed.year}-Q{((parsed.month - 1) // 3) + 1}"


def aggregate_workbook(input_path: Path, output_path: Path) -> dict[str, int]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        raw_header = list(next(rows))
        headers = canonical_headers(raw_header)
        groups: dict[tuple[str, str], FeatureAccumulator] = defaultdict(FeatureAccumulator)
        skipped = 0
        source_rows = 0
        for raw_row in rows:
            values = list(raw_row)
            if not any(value not in (None, "") for value in values):
                skipped += 1
                continue
            if row_looks_like_header(values, raw_header):
                skipped += 1
                continue
            source_rows += 1
            mapped = dict(zip(headers, values))
            period = calendar_period(mapped.get("question_time"))
            code = normalize_stock_code(mapped.get("stock_code")) or "unknown"
            if period is None:
                period = "unknown"
            key = (code, period)
            groups[key].add(mapped)

        fieldnames = [
            "period",
            "stock_code",
            "company_name",
            "question_count",
            "unique_user_count",
            "registered_user_count",
            "browser_user_count",
            "missing_user_type_count",
            "question_length_mean",
            "question_length_min",
            "question_length_max",
            "fiscal_quarter_mode",
            "label_company_violation",
            "label_company_violation_consistent",
            "post_event_replied_count",
            "post_event_unreplied_count",
        ]
        for category in KEYWORDS:
            fieldnames.extend([f"{category}_count", f"{category}_share"])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for (code, period), accumulator in sorted(groups.items(), key=lambda item: item[0]):
                writer.writerow(accumulator.as_row(period))
        return {"source_rows": source_rows, "skipped_rows": skipped, "feature_rows": len(groups)}
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Aggregate Q&A into company-period research features")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(aggregate_workbook(args.input, args.output))


if __name__ == "__main__":
    main()
